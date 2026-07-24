from __future__ import annotations

import html
import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
WORKBOOK = ROOT / "フルーツ送料一覧表.xlsx"
OUT_DIR = ROOT / "scripts" / "shopify_setup"
OFFICIAL_LINE_URL = "https://lin.ee/sEgCPJs"

DESTINATIONS = [
    ("北海道", 4),
    ("北東北", 5),
    ("南東北", 6),
    ("関東", 7),
    ("信越", 8),
    ("北陸", 9),
    ("中部", 10),
    ("関西", 11),
    ("中国", 12),
    ("四国", 13),
    ("九州", 14),
    ("沖縄", 15),
]


def text(value: object) -> str:
    return "" if value is None else str(value).strip()


def yen(value: object) -> str:
    if value is None:
        return "-"
    amount = int(float(value))
    return f"¥{amount:,}"


def prefectures(ws, col: int) -> list[str]:
    values: list[str] = []
    for row in range(5, 13):
        value = text(ws.cell(row, col).value)
        if value:
            values.append(value)
    return values


def read_rates() -> tuple[list[dict[str, object]], list[dict[str, object]]]:
    wb = load_workbook(WORKBOOK, data_only=True, read_only=False)
    ws = wb.active

    areas = [
        {
            "name": name,
            "prefectures": prefectures(ws, col),
        }
        for name, col in DESTINATIONS
    ]

    rows: list[dict[str, object]] = []
    current_origin = ""
    for row in range(13, ws.max_row + 1):
        origin = text(ws.cell(row, 2).value)
        size = ws.cell(row, 3).value
        if origin:
            current_origin = origin
        if not current_origin or size is None:
            continue
        rates = {name: ws.cell(row, col).value for name, col in DESTINATIONS}
        if all(value is None for value in rates.values()):
            continue
        rows.append(
            {
                "origin": current_origin,
                "size": int(float(size)),
                "rates": rates,
            }
        )

    return areas, rows


def build_page_body(areas: list[dict[str, object]], rows: list[dict[str, object]]) -> str:
    parts: list[str] = []
    parts.append('<div class="ship-copy">')
    parts.append(
        "<p>送料は発送元・箱サイズ・お届け先エリアによって異なります。"
        "以下はヤマト運輸宅急便の送料目安です。</p>"
    )
    parts.append(
        "<p>実際の送料は、カートページまたは決済画面でお届け先を入力した際に表示される金額をご確認ください。</p>"
    )
    parts.append("</div>")

    parts.append('<div class="ship-area-block">')
    parts.append("<h2>お届け先エリア</h2>")
    parts.append('<div class="ship-area-grid">')
    for area in areas:
        name = html.escape(str(area["name"]))
        prefs = "・".join(html.escape(pref) for pref in area["prefectures"])
        parts.append(f"<dl><dt>{name}</dt><dd>{prefs}</dd></dl>")
    parts.append("</div>")
    parts.append("</div>")

    parts.append('<div class="ship-rate-block">')
    parts.append("<h2>送料一覧表</h2>")
    parts.append("<p>スマートフォンでは表を横にスクロールしてご確認ください。金額は税込です。</p>")
    parts.append('<div class="ship-table-scroll">')
    parts.append('<table class="ship-rate-table">')
    parts.append("<thead><tr>")
    parts.append('<th scope="col" class="ship-origin">発送元</th>')
    parts.append('<th scope="col" class="ship-size">箱サイズ</th>')
    for name, _ in DESTINATIONS:
        parts.append(f'<th scope="col">{html.escape(name)}</th>')
    parts.append("</tr></thead><tbody>")

    origin_counts: dict[str, int] = {}
    for row in rows:
        origin_counts[str(row["origin"])] = origin_counts.get(str(row["origin"]), 0) + 1

    rendered_origins: set[str] = set()
    for row in rows:
        origin = str(row["origin"])
        parts.append("<tr>")
        if origin not in rendered_origins:
            parts.append(
                f'<th scope="rowgroup" rowspan="{origin_counts[origin]}" class="ship-origin">'
                f"{html.escape(origin)}</th>"
            )
            rendered_origins.add(origin)
        parts.append(f'<th scope="row" class="ship-size">{row["size"]}サイズ</th>')
        rates = row["rates"]
        assert isinstance(rates, dict)
        for name, _ in DESTINATIONS:
            parts.append(f"<td>{yen(rates.get(name))}</td>")
        parts.append("</tr>")

    parts.append("</tbody></table></div>")
    parts.append("</div>")
    return "".join(parts)


def build_shipping_policy_body() -> str:
    return (
        "<h2>送料と配送について</h2>"
        "<p>祝果では、生産者を中心とした全国の提携パートナーから、旬の果物を新鮮なまま直送でお送りしています。</p>"
        "<p>果物本来のおいしさをお届けするため、それぞれに最適なタイミングで収穫・発送を行っており、配送日のご指定は承っておりません。</p>"
        "<p>また、発送元や箱のサイズにより送料が異なりますので、ご注文前に以下の内容をご確認ください。</p>"
        "<h2>配送業者・配送方法</h2>"
        "<p>商品はすべてヤマト運輸の宅急便でお届けいたします。</p>"
        "<p>※配送は日本国内に限らせていただきます。</p>"
        f'<p>ご注文状況の確認や発送通知は、会員登録後にLINE連携いただくとスムーズにご利用いただけます。会員登録は <a href="/account">こちら</a> ／ 公式LINEは <a href="{OFFICIAL_LINE_URL}" target="_blank" rel="noopener">こちら</a>からご利用いただけます。</p>'
        "<h2>発送元</h2>"
        "<p>商品は、生産者を中心とした全国の提携パートナーから直送しております。</p>"
        "<p>ご注文いただいた商品が複数の農家からの発送となる場合、それぞれ別送となります。</p>"
        "<p>発送元や箱サイズにより送料が異なります。</p>"
        "<h2>送料について</h2>"
        "<p>送料は、以下の3つの要素に基づいて決定されます。</p>"
        "<ul>"
        "<li>【発送元】商品を出荷する提携農家などの所在地</li>"
        "<li>【お届け先】ご注文者様の配送先地域</li>"
        "<li>【箱のサイズ】ご注文内容に応じた梱包サイズ（例：60サイズ／80サイズ など）</li>"
        "</ul>"
        "<p>※送料の目安はカートページまたは決済画面でご確認いただけます。</p>"
        '<p>※送料一覧は <a href="/pages/shipping-rates">こちら</a>からご覧いただけます。</p>'
        "<p>※祝果では離島への発送も承っており、追加料金は発生いたしません。</p>"
        "<p>ご注文前に送料を確認されたい場合は、カートページにてお届け先を入力いただくと、送料のお見積もりをご確認いただけます。</p>"
        "<p>ご不明点がありましたら、ページ右下のチャットよりお気軽にお問い合わせください。</p>"
    )


def write_json(path: Path, payload: dict[str, object]) -> None:
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main() -> None:
    areas, rows = read_rates()
    body = build_page_body(areas, rows)
    write_json(
        OUT_DIR / "page_shipping_rates.variables.json",
        {
            "page": {
                "title": "送料一覧",
                "handle": "shipping-rates",
                "body": body,
                "isPublished": True,
                "templateSuffix": "shipping-rates",
            }
        },
    )
    write_json(OUT_DIR / "update_shipping_policy.variables.json", {"body": build_shipping_policy_body()})
    print(json.dumps({"areas": len(areas), "rate_rows": len(rows)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
